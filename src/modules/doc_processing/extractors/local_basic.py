from __future__ import annotations

import csv
import io
import json
import os

from src.modules.doc_processing.models import ExtractedDocument


def _decode_text(raw: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return raw.decode(enc).strip()
        except Exception:
            continue
    return raw.decode("utf-8", errors="ignore").strip()


def _markdown_table(rows: list[list[object]]) -> str:
    cleaned = [["" if c is None else str(c).replace("\n", " ").strip() for c in row] for row in rows]
    cleaned = [row for row in cleaned if any(cell for cell in row)]
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    padded = [row + [""] * (width - len(row)) for row in cleaned]
    header = padded[0]
    sep = ["---"] * width
    body = padded[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(sep) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _extract_csv(data: bytes, filename: str) -> ExtractedDocument:
    text = _decode_text(data)
    reader = csv.reader(io.StringIO(text))
    rows = [row for row in reader]
    table = _markdown_table(rows)
    return ExtractedDocument(
        markdown=f"# {filename}\n\n{table or text or '[empty csv document]'}",
        metadata={"kind": "spreadsheet", "format": "csv", "rows": len(rows)},
    )


def _extract_pdf(data: bytes, filename: str) -> ExtractedDocument:
    try:
        from pypdf import PdfReader  # type: ignore[import-untyped]
    except ImportError as e:
        raise ValueError("pypdf is not installed") from e

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"## Page {idx}\n\n{text}")
    if not pages:
        return ExtractedDocument(
            markdown=f"# {filename}\n\n[No extractable text found. OCR is required for this document.]",
            metadata={"kind": "pdf_scanned", "pages": len(reader.pages), "ocr_required": True},
            warnings=["PDF has no extractable text; OCR is required."],
        )
    return ExtractedDocument(
        markdown=f"# {filename}\n\n" + "\n\n".join(pages),
        metadata={"kind": "pdf_text", "pages": len(reader.pages), "ocr_required": False},
    )


def _extract_docx(data: bytes, filename: str) -> ExtractedDocument:
    try:
        from docx import Document  # type: ignore[import-untyped]
    except ImportError as e:
        raise ValueError("python-docx is not installed") from e

    doc = Document(io.BytesIO(data))
    parts: list[str] = [f"# {filename}"]
    for paragraph in doc.paragraphs:
        text = (paragraph.text or "").strip()
        if text:
            parts.append(text)
    for idx, table in enumerate(doc.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        table_md = _markdown_table(rows)
        if table_md:
            parts.append(f"## Table {idx}\n\n{table_md}")
    return ExtractedDocument(
        markdown="\n\n".join(parts).strip() or f"# {filename}\n\n[empty docx document]",
        metadata={"kind": "document", "format": "docx", "tables": len(doc.tables)},
    )


def _extract_xlsx(data: bytes, filename: str) -> ExtractedDocument:
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
    except ImportError as e:
        raise ValueError("openpyxl is not installed") from e

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = [f"# {filename}"]
    sheet_names: list[str] = []
    total_rows = 0
    for ws in wb.worksheets:
        sheet_names.append(ws.title)
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        total_rows += len(rows)
        table = _markdown_table(rows)
        if table:
            parts.append(f"## Sheet: {ws.title}\n\n{table}")
    return ExtractedDocument(
        markdown="\n\n".join(parts).strip() or f"# {filename}\n\n[empty xlsx document]",
        metadata={"kind": "spreadsheet", "format": "xlsx", "sheets": sheet_names, "rows": total_rows},
    )


def _extract_xls(data: bytes, filename: str) -> ExtractedDocument:
    try:
        import xlrd  # type: ignore[import-untyped]
    except ImportError as e:
        raise ValueError("xlrd is not installed") from e

    book = xlrd.open_workbook(file_contents=data)
    parts = [f"# {filename}"]
    total_rows = 0
    sheet_names: list[str] = []
    for sheet in book.sheets():
        sheet_names.append(sheet.name)
        rows = [[sheet.cell_value(rx, cx) for cx in range(sheet.ncols)] for rx in range(sheet.nrows)]
        total_rows += len(rows)
        table = _markdown_table(rows)
        if table:
            parts.append(f"## Sheet: {sheet.name}\n\n{table}")
    return ExtractedDocument(
        markdown="\n\n".join(parts).strip() or f"# {filename}\n\n[empty xls document]",
        metadata={"kind": "spreadsheet", "format": "xls", "sheets": sheet_names, "rows": total_rows},
    )


def extract_document(data: bytes, filename: str, mime_type: str | None) -> ExtractedDocument:
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in {".txt", ".md"} or "text/" in (mime_type or ""):
        return ExtractedDocument(
            markdown=f"# {filename}\n\n{_decode_text(data) or '[empty text document]'}",
            metadata={"kind": "text", "format": ext.lstrip(".") or mime_type},
        )
    if ext == ".json":
        try:
            obj = json.loads(_decode_text(data))
            body = json.dumps(obj, ensure_ascii=False, indent=2)
        except Exception:
            body = _decode_text(data)
        return ExtractedDocument(
            markdown=f"# {filename}\n\n```json\n{body}\n```",
            metadata={"kind": "structured", "format": "json"},
        )
    if ext == ".csv":
        return _extract_csv(data, filename)
    if ext == ".pdf":
        return _extract_pdf(data, filename)
    if ext == ".docx":
        return _extract_docx(data, filename)
    if ext == ".xlsx":
        return _extract_xlsx(data, filename)
    if ext == ".xls":
        return _extract_xls(data, filename)
    if ext == ".doc":
        raise ValueError("Legacy .doc is not supported yet. Convert the file to .docx.")
    raise ValueError(f"Unsupported document type: {ext or mime_type or filename}")
